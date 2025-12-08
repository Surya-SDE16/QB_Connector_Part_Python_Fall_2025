"""
Item list comparison module with dynamic Excel header detection,
QuickBooks linking, conflict detection, unified 'data_mismatch'
classification, JSON export with full field coverage,
JSON reporting of number of perfectly matching items,
and new items to be added (and pushed to QuickBooks).
"""

from __future__ import annotations
from dataclasses import dataclass
from typing import Optional
from pathlib import Path
from openpyxl import load_workbook
import json
from datetime import datetime, timezone
from typing import Any, Dict

# Import your QuickBooks gateway
from quickbook_connector.qb_gateway import (
    fetch_quickbooks_inventory,
    add_items_to_quickbooks,
)
from quickbook_connector.models import InventoryItem


# ----------------------------------------------------------------------
# DATA MODELS
# ----------------------------------------------------------------------
@dataclass
class Item:
    id: str
    name: str
    price: float


@dataclass
class Conflict:
    id: str
    source1: Optional[Item]  # QB record
    source2: Optional[Item]  # Excel record
    mismatched_fields: list[str]


@dataclass
class ComparisonReport:
    total_source1: int
    total_source2: int
    matching_items: int
    source1_only: list[Item]
    source2_only: list[Item]
    conflicts: list[Conflict]
    add_new_items: list[Item]  # Excel-only items to add to QB


# ----------------------------------------------------------------------
# QUICKBOOKS READER
# ----------------------------------------------------------------------
def fetch_items_from_quickbooks_linked() -> list[Item]:
    qb_items = fetch_quickbooks_inventory()  # returns list of dicts

    items: list[Item] = []
    for entry in qb_items:
        items.append(
            Item(
                id=str(entry["record_id"]),
                name=str(entry["name"]),
                price=float(entry["price"]),
            )
        )
    return items


# ----------------------------------------------------------------------
# EXCEL HEADER MAPPING
# ----------------------------------------------------------------------
def detect_column(headers: list[str], possible_names: list[str]) -> Optional[int]:
    """Return the index of the first matching header, case-insensitive."""
    headers_lower = [h.lower() if h else "" for h in headers]
    for idx, header in enumerate(headers_lower):
        for name in possible_names:
            if header == name.lower():
                return idx
    return None


# ----------------------------------------------------------------------
# ROBUST EXCEL READER
# ----------------------------------------------------------------------
def read_items_from_excel(path: Path) -> list[Item]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h else "" for h in next(rows)]
    except StopIteration:
        return []

    # Detect columns dynamically
    id_idx = detect_column(headers, ["ID", "ItemID", "Item Id"])
    name_idx = detect_column(headers, ["Name", "Item Name", "Description"])
    price_idx = detect_column(headers, ["Price", "Sales Price", "Cost"])

    if id_idx is None or name_idx is None:
        raise ValueError("Could not detect required 'ID' or 'Name' columns in Excel")

    items: list[Item] = []

    for row in rows:
        item_id = row[id_idx] if id_idx < len(row) else None
        name = row[name_idx] if name_idx < len(row) else None
        price = (
            row[price_idx] if price_idx is not None and price_idx < len(row) else 0.0
        )

        if not item_id or not name:
            continue

        try:
            price = float(price) if price not in (None, "") else 0.0
        except (TypeError, ValueError):
            price = 0.0

        items.append(Item(id=str(item_id).strip(), name=str(name).strip(), price=price))

    return items


# ----------------------------------------------------------------------
# COMPARISON LOGIC
# ----------------------------------------------------------------------
def compare_item_lists(list1: list[Item], list2: list[Item]) -> ComparisonReport:
    map1 = {item.id: item for item in list1}  # QuickBooks
    map2 = {item.id: item for item in list2}  # Excel

    matching: list[Item] = []
    only_1: list[Item] = []
    only_2: list[Item] = []
    conflicts: list[Conflict] = []
    new_items: list[Item] = []

    # Compare items for conflicts or matches
    for item_id, item1 in map1.items():
        item2 = map2.get(item_id)
        if item2:
            mismatches: list[str] = []

            if item1.name != item2.name:
                mismatches.append("Name")
            if float(item1.price) != float(item2.price):
                mismatches.append("Price")

            if not mismatches:
                matching.append(item1)
            else:
                conflicts.append(
                    Conflict(
                        id=item_id,
                        source1=item1,
                        source2=item2,
                        mismatched_fields=mismatches,
                    )
                )
        else:
            only_1.append(item1)

    # Detect Excel-only items as new items to be added
    for item_id, item2 in map2.items():
        if item_id not in map1:
            only_2.append(item2)
            new_items.append(item2)

    return ComparisonReport(
        total_source1=len(list1),
        total_source2=len(list2),
        matching_items=len(matching),
        source1_only=only_1,
        source2_only=only_2,
        conflicts=conflicts,
        add_new_items=new_items,
    )


# ----------------------------------------------------------------------
# JSON EXPORT
# ----------------------------------------------------------------------


def write_conflicts_to_json(report: ComparisonReport, path: Path) -> None:
    json_output: Dict[str, Any] = {
        "status": "success",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "same_items": report.matching_items,
        "conflicts": [],
        "add_new_items": [],
        "error": None,
    }

    # 1) Same-ID conflicts
    for c in report.conflicts:
        if c.source1 and c.source2:
            json_output["conflicts"].append(
                {
                    "record_id": c.id,
                    "qb_name": c.source1.name,
                    "excel_name": c.source2.name,
                    "qb_price": c.source1.price,
                    "excel_price": c.source2.price,
                    "reason": "data_mismatch",
                }
            )

    # 2) Items only in QuickBooks
    for item in report.source1_only:
        json_output["conflicts"].append(
            {
                "record_id": item.id,
                "qb_name": item.name,
                "excel_name": None,
                "qb_price": item.price,
                "excel_price": None,
                "reason": "missing_in_excel",
            }
        )

    # 3) New items from Excel
    for item in report.add_new_items:
        json_output["add_new_items"].append(
            {
                "id": item.id,
                "name": item.name,
                "price": item.price,
            }
        )

    with open(path, "w", encoding="utf-8") as f:
        json.dump(json_output, f, indent=4)

    print(f"JSON conflict file created: {path}")


# ----------------------------------------------------------------------
# PUSH NEW ITEMS INTO QUICKBOOKS
# ----------------------------------------------------------------------
def push_new_items_to_quickbooks(report: ComparisonReport) -> None:
    if not report.add_new_items:
        print("No Excel-only items to add to QuickBooks.")
        return

    qb_items: list[InventoryItem] = []
    for item in report.add_new_items:
        qb_items.append(
            InventoryItem(
                record_id=item.id,
                name=item.name,
                price=float(item.price),
                source="excel",
            )
        )

    print(f"\nAdding {len(qb_items)} new Excel items to QuickBooks...")
    add_items_to_quickbooks(qb_items)


# ----------------------------------------------------------------------
# PRETTY PRINT
# ----------------------------------------------------------------------
def print_report(report: ComparisonReport) -> None:
    print("\n=== Item Comparison Report ===")
    print(f"Total Items in QuickBooks: {report.total_source1}")
    print(f"Total Items in Excel: {report.total_source2}")
    print(f"Matching Items: {report.matching_items}")
    print(f"Conflicts: {len(report.conflicts)}")
    print(f"New Items to Add: {len(report.add_new_items)}\n")

    print("Items only in QuickBooks:")
    if report.source1_only:
        for item in report.source1_only:
            print(f" - {item.id}: {item.name} (price={item.price})")
    else:
        print(" - None")

    print("\nItems only in Excel:")
    if report.source2_only:
        for item in report.source2_only:
            print(f" - {item.id}: {item.name} (price={item.price})")
    else:
        print(" - None")

    print("\nConflicts:")
    if report.conflicts:
        for c in report.conflicts:
            print(f" - {c.id}: mismatch in {', '.join(c.mismatched_fields)}")
    else:
        print(" - None")

    print("\nNew Items to Add:")
    if report.add_new_items:
        for item in report.add_new_items:
            print(f" - {item.id}: {item.name} (price={item.price})")
    else:
        print(" - None")


# ----------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------
if __name__ == "__main__":
    excel_path = Path(
        "C:/Users/PotharajuS/QB_Connector_Part_Python_Fall_2025/company_data.xlsx"
    )

    qb_items = fetch_items_from_quickbooks_linked()
    excel_items = read_items_from_excel(excel_path)

    report = compare_item_lists(qb_items, excel_items)
    print_report(report)

    write_conflicts_to_json(report, Path("conflicts_output.json"))
    push_new_items_to_quickbooks(report)
